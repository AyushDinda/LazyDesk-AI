import os
import pymupdf
import docx
from PIL import Image, ImageDraw

def create_sample_messy_folder(target_dir: str):
    os.makedirs(target_dir, exist_ok=True)

    # 1. Financial Text
    with open(os.path.join(target_dir, "INV-98234-billing-oct.txt"), "w", encoding="utf-8") as f:
        f.write(
            "ACME Cloud Services Inc.\n"
            "Invoice #: INV-98234\n"
            "Date: October 15, 2024\n"
            "Customer: LazyDesk Dev Team\n"
            "Description: High-Performance GPU Cloud Compute (72 hours)\n"
            "Subtotal: $240.00\n"
            "Tax: $19.20\n"
            "Total Amount Due: $259.20\n"
            "Payment Status: Paid via Credit Card ending in 4092\n"
        )

    # 2. Financial CSV
    with open(os.path.join(target_dir, "q3_expenses_records.csv"), "w", encoding="utf-8") as f:
        f.write(
            "Date,Category,Description,Amount,Status\n"
            "2024-09-01,Hosting,Vercel Pro Tier,20.00,Approved\n"
            "2024-09-05,Domain,Namecheap DNS,12.50,Approved\n"
            "2024-09-15,API,Ollama Server Infrastructure,45.00,Approved\n"
            "2024-09-30,Hardware,NVMe SSD Backup Drive,120.00,Approved\n"
        )

    # 3. Academic PDF (Syllabus)
    pdf_path_1 = os.path.join(target_dir, "CS204_Fall2024_Syllabus.pdf")
    doc1 = pymupdf.open()
    page1 = doc1.new_page()
    page1.insert_text(
        (50, 72),
        "Department of Computer Science - University of Technology\n"
        "CS 204: Data Structures & Algorithms\n"
        "Instructor: Dr. Alan Turing (Office: Turing Hall 302)\n\n"
        "Course Overview:\n"
        "This course covers asymptotic complexity, trees, graphs, dynamic programming,\n"
        "and algorithmic paradigms with weekly programming assignments.\n\n"
        "Grading Rubric:\n"
        "- Homework & Assignments: 30%\n"
        "- Midterm Exam: 30%\n"
        "- Final Term Project: 40%\n",
        fontsize=12
    )
    doc1.save(pdf_path_1)
    doc1.close()

    # 4. Technical Spec PDF
    pdf_path_2 = os.path.join(target_dir, "distributed_cache_design_doc.pdf")
    doc2 = pymupdf.open()
    page2 = doc2.new_page()
    page2.insert_text(
        (50, 72),
        "Engineering RFC: Distributed In-Memory Cache Layer\n"
        "Author: Infrastructure Team\n"
        "Status: Approved for Sprint 12\n\n"
        "Architecture Overview:\n"
        "We propose deploying a consistent hashing distributed cache with redis-compatible protocol.\n"
        "Key requirements: sub-millisecond p99 read latency and automatic failover handling.\n"
        "Replication factor: 3 nodes per cluster shard.\n",
        fontsize=12
    )
    doc2.save(pdf_path_2)
    doc2.close()

    # 5. Work Document DOCX (Contract / Agreement)
    docx_path = os.path.join(target_dir, "Freelance_Consulting_Agreement_Draft.docx")
    d = docx.Document()
    d.add_heading("Independent Contractor Services Agreement", level=1)
    d.add_paragraph("This Agreement is entered into between NovaTech Labs ('Client') and Ayush ('Contractor').")
    d.add_paragraph("Scope of Work: Architecture design, backend FastAPI engineering, and local AI pipeline development.")
    d.add_paragraph("Term: 3 months commencing October 1, 2024. Compensation rate: Fixed milestone disbursements.")
    d.save(docx_path)

    # 6. Notes Markdown
    with open(os.path.join(target_dir, "weekend_ideas_and_todo.md"), "w", encoding="utf-8") as f:
        f.write(
            "# Weekly Brainstorming & Personal Notes\n\n"
            "- [x] Research local LLM quantization (GGUF, Q4_K_M)\n"
            "- [x] Test PyMuPDF text streaming performance on 100-page documents\n"
            "- [ ] Implement undo transactional safeguards in SQLite\n"
            "- [ ] Review Hacktoberfest Week 3 criteria\n"
            "- [ ] Prepare sample demonstration video\n"
        )

    # 7. Image Asset PNG
    img_path = os.path.join(target_dir, "screenshot_dashboard_v1.png")
    img = Image.new("RGB", (640, 480), color=(34, 75, 55))
    draw = ImageDraw.Draw(img)
    draw.rectangle([40, 40, 600, 440], fill=(245, 242, 235), outline=(27, 67, 50), width=3)
    draw.text((60, 60), "LazyDesk AI Mockup Preview", fill=(27, 67, 50))
    img.save(img_path)

    print(f"Sample messy folder successfully seeded at: {target_dir}")

if __name__ == "__main__":
    current_workspace = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    target = os.path.join(current_workspace, "sample_messy_folder")
    create_sample_messy_folder(target)
